"""Document text extraction: PDF via pdfplumber, DOCX via python-docx."""

from pathlib import Path

import pdfplumber
from docx import Document as DocxDocument


class ParsingError(Exception):
    pass


def extract_pdf_text(path: Path) -> str:
    try:
        with pdfplumber.open(path) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
    except Exception as error:  # noqa: BLE001 - normalize parser errors
        raise ParsingError(f"Failed to read PDF: {error}") from error
    return "\n\n".join(page.strip() for page in pages if page.strip())


def extract_docx_text(path: Path) -> str:
    try:
        document = DocxDocument(path)
    except Exception as error:  # noqa: BLE001
        raise ParsingError(f"Failed to read DOCX: {error}") from error
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells if cell.text.strip())
    return "\n".join(parts)


def extract_text(path: Path) -> str:
    """Extract plain text from a CV file (PDF or DOCX) by extension."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_text(path)
    if suffix in {".docx", ".doc"}:
        return extract_docx_text(path)
    raise ParsingError(f"Unsupported file type: {suffix}")
